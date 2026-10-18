"""Exercise the Rust CLI and real local engines, without opening the GUI.

Build first: cargo build -p orbit-core --bin orbit-convert
Run: python scripts/smoke.py
"""
from __future__ import annotations
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
from reportlab.lib import colors

root=Path(__file__).resolve().parents[1]
binary=root/'target'/'debug'/('orbit-convert.exe' if sys.platform=='win32' else 'orbit-convert')
if not binary.is_file(): raise SystemExit('Build the Rust CLI first: cargo build -p orbit-core --bin orbit-convert')
environment=os.environ.copy()
environment.setdefault('ORBIT_ENGINE_PYTHON',sys.executable)
environment.setdefault('ORBIT_ENGINE_SCRIPT',str(root/'engine'/'worker.py'))
office=environment.get('ORBIT_OFFICE') or shutil.which('soffice')
for candidate in ['/Applications/LibreOffice.app/Contents/MacOS/soffice',r'C:\Program Files\LibreOffice\program\soffice.com']:
    if not office and Path(candidate).is_file(): office=candidate
if office: environment['ORBIT_OFFICE']=office
if environment.get('ORBIT_REQUIRE_OFFICE') == '1' and not office:
    raise SystemExit('LibreOffice is required for this native validation run but was not detected')
with tempfile.TemporaryDirectory(prefix='orbit-smoke-') as directory:
    folder=Path(directory)
    def run(source:Path,fmt:str)->dict:
        request=folder/'request.json'
        request.write_text(json.dumps({'input':str(source),'output_dir':str(folder),'format':fmt}),encoding='utf-8')
        before=source.read_bytes()
        result=subprocess.run([str(binary),str(request)],env=environment,capture_output=True,text=True)
        if result.returncode:
            raise RuntimeError(result.stderr)
        receipt=json.loads(result.stdout)
        assert source.read_bytes()==before,'Original was modified'
        assert receipt['files'] and all(Path(p).stat().st_size>0 for p in receipt['files'])
        print(f'PASS {source.suffix} -> {fmt}',flush=True)
        return receipt
    image=folder/'image.png';Image.new('RGBA',(120,80),(0,180,100,128)).save(image)
    run(image,'jpg');run(image,'webp');run(image,'tiff');run(image,'bmp')
    pdf=Path(run(image,'pdf')['files'][0]);run(pdf,'png')
    table_pdf=folder/'table.pdf'
    table=Table([['Product','Quantity'],['Apples','12'],['Oranges','8']],colWidths=[130,100])
    table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),1,colors.black)]))
    SimpleDocTemplate(str(table_pdf)).build([table])
    for fmt in ['docx','xlsx','csv','txt','jpg']:
        run(table_pdf,fmt)
    if office:
        doc=Document();doc.add_heading('Orbit smoke test',0);doc.add_paragraph('Local Word to PDF conversion.');source=folder/'word.docx';doc.save(source)
        word_pdf=Path(run(source,'pdf')['files'][0]);run(word_pdf,'docx');run(word_pdf,'txt');run(source,'odt');run(source,'rtf');run(source,'txt');run(source,'html')
        book=Workbook();sheet=book.active;sheet.append(['Product','Quantity']);sheet.append(['Apples',12]);sheet.append(['Oranges',8]);source=folder/'excel.xlsx';book.save(source)
        excel_pdf=Path(run(source,'pdf')['files'][0]);receipt=run(excel_pdf,'xlsx');run(excel_pdf,'csv')
        assert 'Apples' in str(list(load_workbook(receipt['files'][0]).active.values))
        run(source,'csv');run(source,'tsv');run(source,'ods')
    else: print('SKIP Office routes: LibreOffice not installed')
print('Smoke checks complete.')
