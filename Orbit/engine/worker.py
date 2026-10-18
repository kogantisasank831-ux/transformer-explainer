"""Offline PDF/OCR worker. The Rust host owns jobs, cancellation and publication.

PDF -> Word reconstructs content; it does not promise the original layout.
PDF -> Excel exports detected table values, never the original formulas.
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import logging
import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from pydantic import BaseModel, ConfigDict, Field, field_validator

LOGGER = logging.getLogger("orbit.engine")
MAX_PAGES = 300
MAX_PIXELS = 40_000_000


class Options(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)
    quality: int = Field(default=90, ge=1, le=100)
    dpi: int = Field(default=150, ge=72, le=300)
    pages: str = Field(default="", max_length=256)
    ocr: bool = False
    language: str = Field(default="eng", pattern=r"^[A-Za-z0-9_+]{1,64}$")
    docx_mode: str = Field(default="editable", pattern=r"^(editable|page_images)$")
    single_page_sheets: bool = False


class Request(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)
    input: Path
    output_dir: Path
    format: str
    options: Options = Field(default_factory=Options)
    tesseract: Path | None = None

    @field_validator("format")
    @classmethod
    def known_format(cls, value: str) -> str:
        if value not in {"pdf", "docx", "xlsx", "csv", "jpg", "png", "txt"}:
            raise ValueError("Unsupported worker output format")
        return value


class Response(BaseModel):
    files: list[str]
    warnings: list[str] = Field(default_factory=list)


@dataclass(frozen=True)
class Word:
    text: str
    x0: float
    top: float
    x1: float
    bottom: float


def selected_pages(spec: str, total: int) -> list[int]:
    """Validate one-based ranges and return unique, document-ordered indices."""
    if total > MAX_PAGES:
        raise ValueError(f"PDFs are limited to {MAX_PAGES} pages per job")
    if not spec.strip():
        return list(range(total))
    chosen: set[int] = set()
    for segment in spec.split(","):
        match = re.fullmatch(r"\s*(\d+)\s*(?:-\s*(\d+)\s*)?", segment)
        if not match:
            raise ValueError("Pages must look like 1-3,5")
        first = int(match[1])
        last = int(match[2] or match[1])
        if not 1 <= first <= last <= total:
            raise ValueError(f"Page range must be between 1 and {total}")
        chosen.update(range(first - 1, last))
    return sorted(chosen)


def render_page(pdf: Any, index: int, dpi: int) -> Any:
    page = pdf[index]
    try:
        width, height = page.get_size()
        scale = dpi / 72
        if width * height * scale * scale > MAX_PIXELS:
            raise ValueError("Rendered page exceeds 40 megapixels; reduce DPI")
        bitmap = page.render(scale=scale)
        try:
            return bitmap.to_pil().convert("RGB").copy()
        finally:
            bitmap.close()
    finally:
        page.close()


def tesseract_path(request: Request) -> str:
    binary = str(request.tesseract) if request.tesseract else shutil.which("tesseract")
    if not binary or not Path(binary).is_file():
        raise ValueError("OCR requires Tesseract. Install it and its selected language data, then restart Orbit.")
    return binary


def ocr_words(image: Any, request: Request) -> list[Word]:
    with tempfile.TemporaryDirectory(prefix="orbit-ocr-") as directory:
        image_path = Path(directory) / "page.png"
        image.save(image_path)
        # Inherit the worker process group so the Rust host can kill the full tree.
        result = subprocess.run(
            [tesseract_path(request), str(image_path), "stdout", "-l", request.options.language, "tsv"],
            stdin=subprocess.DEVNULL, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
            text=True, encoding="utf-8", errors="replace", timeout=180, check=False,
            creationflags=subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0,
        )
        if result.returncode:
            raise ValueError("OCR failed. Check language data and image quality. " + result.stderr[-1000:])
        words: list[Word] = []
        for row in csv.DictReader(io.StringIO(result.stdout), delimiter="\t"):
            text = (row.get("text") or "").strip()
            if not text:
                continue
            x, y = float(row["left"]), float(row["top"])
            words.append(Word(text, x, y, x + float(row["width"]), y + float(row["height"])))
        return words


def group_lines(words: list[Word]) -> list[list[Word]]:
    lines: list[list[Word]] = []
    for word in sorted(words, key=lambda w: (w.top, w.x0)):
        height = max(word.bottom - word.top, 1)
        if lines and abs(word.top - lines[-1][0].top) < max(3, height * .5):
            lines[-1].append(word)
        else:
            lines.append([word])
    return [sorted(line, key=lambda w: w.x0) for line in lines]


def ocr_table(words: list[Word]) -> list[list[str]]:
    """Conservative whitespace table heuristic; always flag it for review."""
    rows: list[list[str]] = []
    for line in group_lines(words):
        cells: list[str] = []
        previous: Word | None = None
        for word in line:
            gap = word.x0 - previous.x1 if previous else 0
            if previous and gap > max(24, (word.bottom - word.top) * 1.8):
                cells.append(word.text)
            elif cells:
                cells[-1] += " " + word.text
            else:
                cells.append(word.text)
            previous = word
        rows.append(cells)
    if sum(len(row) >= 2 for row in rows) < 2:
        return []
    return rows


@dataclass(frozen=True)
class InferredTable:
    bbox: tuple[float, float, float, float]
    data: list[list[str]]

    def extract(self) -> list[list[str]]:
        return self.data


def detect_tables(page: Any) -> list[Any]:
    ruled = [table for table in page.find_tables() if len(table.rows) >= 2 and len(table.columns) >= 2]
    if ruled:
        return ruled
    # Unruled office exports often have right-aligned numbers and left-aligned
    # headings. Pure vertical-edge clustering can truncate those headings.
    words = [Word(w['text'],float(w['x0']),float(w['top']),float(w['x1']),float(w['bottom'])) for w in page.extract_words()]
    inferred: list[InferredTable] = []
    group: list[tuple[list[Word], list[str]]] = []
    def flush() -> None:
        if len(group) >= 2:
            all_words = [w for line,_ in group for w in line]
            inferred.append(InferredTable((min(w.x0 for w in all_words),min(w.top for w in all_words),
                max(w.x1 for w in all_words),max(w.bottom for w in all_words)),[cells for _,cells in group]))
        group.clear()
    for line in group_lines(words):
        cells: list[str] = []
        previous: Word | None = None
        for word in line:
            if previous and word.x0-previous.x1 > max(10,(word.bottom-word.top)*.8):
                cells.append(word.text)
            elif cells:
                cells[-1] += ' '+word.text
            else:
                cells.append(word.text)
            previous=word
        if len(cells)<2:
            flush()
            continue
        if group and (len(cells)!=len(group[-1][1]) or line[0].top-group[-1][0][0].top > 35):
            flush()
        group.append((line,cells))
    flush()
    if inferred:
        return inferred
    return [table for table in page.find_tables({'vertical_strategy':'text','horizontal_strategy':'text',
        'min_words_vertical':3,'min_words_horizontal':1}) if len(table.rows)>=2 and len(table.columns)>=2]


def safe_cell(value: Any) -> str:
    # Extracted document text must not become spreadsheet formulas.
    text = str(value or "").replace("\x00", "")
    text = re.sub(r"[\x01-\x08\x0b\x0c\x0e-\x1f]", "", text)
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) else text


def convert(request: Request) -> Response:
    import pdfplumber
    import pypdfium2 as pdfium
    from PIL import Image, ImageOps

    Image.MAX_IMAGE_PIXELS = 50_000_000
    input_path = request.input.resolve(strict=True)
    if not input_path.is_file() or input_path.stat().st_size > 250 * 1024 * 1024:
        raise ValueError("Choose a file no larger than 250 MiB")
    out = request.output_dir.resolve(strict=True)
    if not out.is_dir() or any(out.iterdir()):
        raise ValueError("Worker output must be a new, empty directory")
    warnings: list[str] = []
    files: list[Path] = []
    ext = input_path.suffix.lower()
    if ext != ".pdf":
        with Image.open(input_path) as opened:
            image = ImageOps.exif_transpose(opened).convert("RGB")
            if image.width * image.height > 50_000_000:
                raise ValueError("Images are limited to 50 megapixels")
            if request.format == "pdf":
                target = out / "converted.pdf"
                image.save(target, "PDF", resolution=request.options.dpi)
                files.append(target)
            elif request.format == "txt":
                text = "\n".join(" ".join(w.text for w in line) for line in group_lines(ocr_words(image, request)))
                if not text.strip():
                    raise ValueError("OCR did not recognise any text")
                target = out / "converted.txt"
                target.write_text(text, encoding="utf-8")
                files.append(target)
                warnings.append("OCR text needs review, especially numbers, names and equations.")
            else:
                raise ValueError("This image conversion is not handled by the PDF worker")
        return Response(files=[str(p) for p in files], warnings=warnings)

    with pdfium.PdfDocument(input_path) as rendered, pdfplumber.open(input_path) as pdf:
        indices = selected_pages(request.options.pages, len(pdf.pages))
        if not indices:
            raise ValueError("The PDF has no pages")
        if request.format in {"jpg", "png"}:
            for index in indices:
                image = render_page(rendered, index, request.options.dpi)
                target = out / f"page-{index+1:04d}.{request.format}"
                image.save(target, quality=request.options.quality) if request.format == "jpg" else image.save(target)
                files.append(target)
        elif request.format == "txt":
            texts: list[str] = []
            for index in indices:
                text = pdf.pages[index].extract_text(layout=False) or ""
                if not text.strip() and request.options.ocr:
                    words = ocr_words(render_page(rendered, index, request.options.dpi), request)
                    text = "\n".join(" ".join(w.text for w in line) for line in group_lines(words))
                    warnings.append(f"Page {index+1}: text recognised with OCR; review it.")
                if not text.strip():
                    warnings.append(f"Page {index+1}: no text detected. Enable OCR for scanned pages.")
                texts.append(text)
            if not any(t.strip() for t in texts):
                raise ValueError("No text extracted. Enable OCR if this is a scanned PDF.")
            target = out / "converted.txt"
            target.write_text("\n\n\f\n\n".join(texts), encoding="utf-8")
            files.append(target)
        elif request.format in {"xlsx", "csv"}:
            from openpyxl import Workbook
            from openpyxl.styles import Font
            workbook = Workbook()
            workbook.remove(workbook.active)
            table_count = 0
            for index in indices:
                page = pdf.pages[index]
                detected = [t.extract() for t in detect_tables(page)]
                if not detected and request.options.ocr:
                    rows = ocr_table(ocr_words(render_page(rendered, index, request.options.dpi), request))
                    if rows:
                        detected = [rows]
                        warnings.append(f"Page {index+1}: OCR inferred columns; review every row.")
                if not detected:
                    warnings.append(f"Page {index+1}: no table detected.")
                for table_index, table in enumerate(detected, 1):
                    rows = [[safe_cell(cell) for cell in row] for row in table if any(str(cell or "").strip() for cell in row)]
                    if not rows:
                        continue
                    table_count += 1
                    title = f"P{index+1}_Table{table_index}"
                    if request.format == "csv":
                        target = out / f"{title}.csv"
                        with target.open("w", encoding="utf-8-sig", newline="") as handle:
                            csv.writer(handle).writerows(rows)
                        files.append(target)
                    else:
                        sheet = workbook.create_sheet(title)
                        for row in rows:
                            sheet.append(row)
                        sheet.freeze_panes = "A2"
                        for cell in sheet[1]:
                            cell.font = Font(bold=True)
                        for column in sheet.columns:
                            sheet.column_dimensions[column[0].column_letter].width = min(60, max(12, max(len(str(c.value or "")) for c in column) + 2))
            if table_count == 0:
                raise ValueError("No tables found. Try OCR for scanned pages, or export text instead.")
            if request.format == "xlsx":
                target = out / "tables.xlsx"
                workbook.save(target)
                files.append(target)
            warnings.append("Extracted values only: formulas, macros and original workbook structure cannot be recovered. Tables need review.")
        elif request.format == "docx":
            from docx import Document
            from docx.shared import Inches, Pt
            document = Document()
            document.styles["Normal"].font.name = "Arial"
            document.styles["Normal"].font.size = Pt(11)
            content_count = 0
            for position, index in enumerate(indices):
                page = pdf.pages[index]
                if position:
                    document.add_page_break()
                section = document.sections[0]
                section.page_width = Pt(float(page.width))
                section.page_height = Pt(float(page.height))
                section.left_margin = section.right_margin = Inches(.55)
                if request.options.docx_mode == "page_images":
                    image = render_page(rendered, index, request.options.dpi)
                    buffer = io.BytesIO()
                    image.save(buffer, format="PNG")
                    buffer.seek(0)
                    available_width = float(page.width) - 79.2
                    available_height = float(page.height) - 144
                    width = min(available_width, available_height * image.width / image.height)
                    document.add_picture(buffer, width=Pt(max(36, width)))
                    content_count += 1
                    continue
                tables = detect_tables(page)
                bboxes = [table.bbox for table in tables]
                def outside_tables(obj: dict[str, Any]) -> bool:
                    x = (float(obj.get("x0", 0)) + float(obj.get("x1", 0))) / 2
                    y = (float(obj.get("top", 0)) + float(obj.get("bottom", 0))) / 2
                    return not any(x0 <= x <= x1 and top <= y <= bottom for x0, top, x1, bottom in bboxes)
                text_page = page.filter(outside_tables)
                blocks: list[tuple[float, str, Any]] = []
                for line in text_page.extract_text_lines():
                    blocks.append((float(line["top"]), "text", line["text"]))
                for table in tables:
                    blocks.append((float(table.bbox[1]), "table", table.extract()))
                if not blocks and request.options.ocr:
                    words = ocr_words(render_page(rendered, index, request.options.dpi), request)
                    blocks = [(line[0].top, "text", " ".join(w.text for w in line)) for line in group_lines(words)]
                    warnings.append(f"Page {index+1}: reconstructed from OCR; review text and layout.")
                if not blocks:
                    warnings.append(f"Page {index+1}: no editable content. Try OCR or page-image mode.")
                for _, kind, value in sorted(blocks, key=lambda b: b[0]):
                    if kind == "text":
                        document.add_paragraph(str(value))
                        content_count += 1
                    else:
                        if not value:
                            continue
                        table_doc = document.add_table(rows=0, cols=max(len(row) for row in value))
                        table_doc.style = "Table Grid"
                        for row in value:
                            cells = table_doc.add_row().cells
                            for col, text in enumerate(row):
                                cells[col].text = str(text or "")
                        content_count += 1
            if not content_count:
                raise ValueError("No editable content found. Enable OCR or use page-image mode.")
            target = out / "converted.docx"
            document.save(target)
            files.append(target)
            warnings.append("Page-image mode embeds pictures and is not editable." if request.options.docx_mode == "page_images" else "Reconstructed text and tables: fonts, columns, drawings and exact layout are not preserved. Embedded pictures are omitted; use page-image mode to retain page appearance.")
        else:
            raise ValueError("Unsupported PDF output format")
    return Response(files=[str(p) for p in files], warnings=list(dict.fromkeys(warnings)))


def doctor() -> dict[str, Any]:
    import importlib
    dependencies = {"pdfplumber": "pdfplumber", "pdfium": "pypdfium2", "word": "docx", "excel": "openpyxl", "images": "PIL"}
    failures: list[str] = []
    for label, module in dependencies.items():
        try:
            importlib.import_module(module)
        except ImportError:
            failures.append(label)
    return {"ready": not failures, "missing": failures, "ocr": bool(shutil.which("tesseract"))}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--request", type=Path)
    parser.add_argument("--response", type=Path, required=True)
    parser.add_argument("--doctor", action="store_true")
    args = parser.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s: %(message)s", stream=sys.stderr)
    try:
        if args.doctor:
            args.response.write_text(json.dumps(doctor()), encoding="utf-8")
        else:
            if not args.request or args.request.stat().st_size > 64 * 1024:
                raise ValueError("Missing or oversized worker request")
            request = Request.model_validate_json(args.request.read_text(encoding="utf-8"))
            response = convert(request)
            args.response.write_text(response.model_dump_json(), encoding="utf-8")
        return 0
    except Exception as error:
        LOGGER.error("%s", error)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
