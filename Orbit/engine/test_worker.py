from __future__ import annotations
import csv
from pathlib import Path
import shutil
import sys

import pytest
from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

sys.path.insert(0, str(Path(__file__).parent))
from worker import Request, Options, convert, selected_pages, safe_cell


@pytest.fixture
def table_pdf(tmp_path: Path) -> Path:
    path=tmp_path/'report.pdf'
    table=Table([['Product','Quantity','Value'],['Apples','12','240'],['Oranges','8','320']], colWidths=[130]*3)
    table.setStyle(TableStyle([('GRID',(0,0),(-1,-1),1,colors.black),('BACKGROUND',(0,0),(-1,0),colors.lightgrey)]))
    SimpleDocTemplate(str(path),pagesize=A4).build([table])
    return path


def request(tmp_path: Path, source: Path, fmt: str, **options: object) -> Request:
    out=tmp_path/f'out-{fmt}-{len(list(tmp_path.iterdir()))}'
    out.mkdir()
    return Request(input=source, output_dir=out, format=fmt, options=Options(**options))


def test_page_ranges_are_bounded_and_document_ordered() -> None:
    assert selected_pages('3,1-2,2',4)==[0,1,2]
    for spec in ['0','2-1','1,','999999999999','x']:
        with pytest.raises(ValueError): selected_pages(spec,4)
    with pytest.raises(ValueError): selected_pages('',301)


def test_formula_like_text_is_neutralised() -> None:
    assert safe_cell('=HYPERLINK("x")').startswith("'")
    assert safe_cell('  +cmd').startswith("'")
    assert safe_cell('normal')=='normal'


def test_pdf_to_excel_has_actual_cells(tmp_path: Path, table_pdf: Path) -> None:
    response=convert(request(tmp_path,table_pdf,'xlsx'))
    book=load_workbook(response.files[0],data_only=False)
    assert book.active['A1'].value=='Product'
    assert book.active['A2'].value=='Apples'
    assert book.active['C3'].value=='320'
    assert any('formulas' in warning for warning in response.warnings)


def test_pdf_to_csv(tmp_path: Path, table_pdf: Path) -> None:
    response=convert(request(tmp_path,table_pdf,'csv'))
    with Path(response.files[0]).open(encoding='utf-8-sig') as handle:
        assert list(csv.reader(handle))[1]==['Apples','12','240']


def test_unruled_table_with_right_aligned_numbers(tmp_path: Path) -> None:
    source=tmp_path/'unruled.pdf'
    canvas=Canvas(str(source));canvas.setFont('Helvetica',11)
    canvas.drawString(55,760,'Product');canvas.drawString(104,760,'Quantity')
    canvas.drawString(55,745,'Apples');canvas.drawRightString(151,745,'12')
    canvas.drawString(55,730,'Oranges');canvas.drawRightString(151,730,'8')
    canvas.save()
    response=convert(request(tmp_path,source,'xlsx'))
    sheet=load_workbook(response.files[0]).active
    assert sheet['B1'].value=='Quantity'
    assert sheet['B3'].value=='8'


def test_pdf_to_editable_word(tmp_path: Path, table_pdf: Path) -> None:
    response=convert(request(tmp_path,table_pdf,'docx'))
    doc=Document(response.files[0])
    assert doc.tables[0].cell(1,0).text=='Apples'
    assert any('pictures are omitted' in warning for warning in response.warnings)


def test_word_page_image_mode_is_explicit(tmp_path: Path, table_pdf: Path) -> None:
    response=convert(request(tmp_path,table_pdf,'docx',docx_mode='page_images'))
    assert len(Document(response.files[0]).inline_shapes)==1
    assert 'not editable' in response.warnings[0]


@pytest.mark.parametrize('fmt',['txt','jpg','png'])
def test_pdf_text_and_rendering(tmp_path: Path, table_pdf: Path, fmt: str) -> None:
    response=convert(request(tmp_path,table_pdf,fmt))
    path=Path(response.files[0])
    if fmt=='txt': assert 'Apples' in path.read_text()
    else:
        with Image.open(path) as image: assert image.width>1000 and image.height>1000


def test_no_table_is_an_error_instead_of_fake_excel(tmp_path: Path) -> None:
    source=tmp_path/'blank.pdf'
    canvas=Canvas(str(source));canvas.showPage();canvas.save()
    with pytest.raises(ValueError,match='No tables'):
        convert(request(tmp_path,source,'xlsx'))


def test_image_to_pdf(tmp_path: Path) -> None:
    source=tmp_path/'photo.png'
    Image.new('RGB',(60,40),'red').save(source)
    response=convert(request(tmp_path,source,'pdf'))
    assert Path(response.files[0]).read_bytes().startswith(b'%PDF')


def test_scanned_pdf_needs_ocr_and_recognises_printed_text(tmp_path: Path) -> None:
    if not shutil.which('tesseract'):
        pytest.skip('Tesseract not installed')
    image=Image.new('RGB',(1600,500),'white')
    drawing=ImageDraw.Draw(image)
    candidates=[Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'),
                Path('C:/Windows/Fonts/arial.ttf'),
                Path('/System/Library/Fonts/Supplemental/Arial.ttf')]
    font_path=next((path for path in candidates if path.is_file()), None)
    if font_path is None:
        pytest.skip('Test font unavailable on this platform')
    drawing.text((80,120),'ORBIT LOCAL CONVERSION 2026',fill='black',font=ImageFont.truetype(str(font_path),54))
    source=tmp_path/'scanned.pdf';image.save(source,'PDF',resolution=150)
    with pytest.raises(ValueError,match='No text'):
        convert(request(tmp_path,source,'txt'))
    response=convert(request(tmp_path,source,'txt',ocr=True,dpi=150))
    assert 'ORBIT' in Path(response.files[0]).read_text()


def test_worker_rejects_nonempty_output(tmp_path: Path, table_pdf: Path) -> None:
    req=request(tmp_path,table_pdf,'txt')
    (req.output_dir/'keep.txt').write_text('keep')
    with pytest.raises(ValueError,match='empty directory'): convert(req)
    assert (req.output_dir/'keep.txt').read_text()=='keep'

